#!/usr/bin/env python3
"""Convierte la tabla principal del Excel (hoja CABLEADO, Tabla_MAIN) a data.json.

Uso:
    pip install openpyxl
    python tools/xlsx_a_json.py "CABLEADO CONEXIONADO Y TIMBRADO BW04.xlsx" data.json

Luego sube data.json a GitHub (reemplaza el anterior) y la app lo cargará sola.
"""
import sys, json, datetime
import openpyxl
from openpyxl.utils import get_column_letter

TAG_HEADER = "CABLE_NUMBER"
ERRORES = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}


def fmt(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(round(v, 4))
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    s = str(v).strip()
    return "" if s in ERRORES else s


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = None
    # 1) hoja que contiene la tabla llamada Tabla_MAIN / Main; 2) hoja con encabezado CABLE_NUMBER
    for s in wb.worksheets:
        if any(n.lower().replace("_", "").replace(" ", "") in ("tablamain", "main") for n in s.tables):
            ws = s
            break
    if ws is None:
        for s in wb.worksheets:
            if TAG_HEADER in [c.value for c in s[1]]:
                ws = s
                break
    if ws is None:
        sys.exit("No encontré la tabla principal (encabezado CABLE_NUMBER en fila 1).")

    header = [fmt(c.value) for c in ws[1]]
    last = max(i for i, h in enumerate(header) if h) + 1
    header = header[:last]
    tag_idx = header.index(TAG_HEADER)

    rows = []
    for r in ws.iter_rows(min_row=2, max_col=last, values_only=True):
        vals = [fmt(v) for v in r]
        if vals[tag_idx]:
            rows.append(vals)

    out = {
        "generado": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "origen": src.split("/")[-1].split("\\")[-1],
        "hoja": ws.title,
        "tagCol": TAG_HEADER,
        "columns": [{"letter": get_column_letter(i + 1), "name": h} for i, h in enumerate(header)],
        "rows": rows,
    }
    with open(dst, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    print(f"OK: {len(rows)} cables, {len(header)} columnas -> {dst}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
